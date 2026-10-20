from flask_babel import gettext as _, lazy_gettext as _l
from io import BytesIO
import re
import tempfile
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile
from datetime import date, datetime, time, timedelta, timezone
from zoneinfo import ZoneInfo
from flask import Blueprint, abort, current_app, flash, g, redirect, render_template, request, send_file, session, url_for
from sqlalchemy import and_, select, func, or_
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.orm import selectinload
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from .models import db, User, AdminPermission, Group, Department, Subject, Event, Booking, LoginGate, AuditLog, SiteSettings, TutorialVideo, utcnow
from .auth import roles_required, admin_permissions_required, normalize_login, hash_password, name_parts, check_password_distinct, lock_gate, clear_gate, audit
from .permissions import ADMIN_PERMISSION_CODES, ADMIN_PERMISSION_CODE_SET, has_admin_permission
from .services import integer, cancel_event, cancel_booking
from .lifecycle import trash_users, restore_users
from .notifications import notify
from .tutorials import clear_rutube_availability_cache, normalize_rutube_url

bp = Blueprint('admin', __name__, url_prefix='/admin')
MOSCOW = ZoneInfo('Europe/Moscow')
MAX_STUDENT_IMPORT = 1500
MAX_XLSX_UNCOMPRESSED = 20 * 1024 * 1024


class StudentImportError(ValueError):
    pass


def audit_line(entry):
    def field(value):
        return ' '.join(str(value or '').split()) or '—'

    actor = f'{field(entry.actor.login)} ({field(entry.actor.full_name)})' if entry.actor else 'Система'
    stamp = entry.created_at.astimezone(MOSCOW).strftime('%d.%m.%Y %H:%M:%S')
    label = current_app.jinja_env.filters['action_label'](entry.action)
    action = f'{field(label)} [{entry.action}]' if label != entry.action else entry.action
    return f'#{entry.id} [{stamp} МСК] {actor} | {action} | {field(entry.entity)} | {field(entry.details)}'


def audit_dates(data):
    try:
        start = date.fromisoformat(data.get('from_date', ''))
        end = date.fromisoformat(data.get('to_date', ''))
    except (TypeError, ValueError):
        raise ValueError(_('Укажите корректные даты начала и окончания.')) from None
    if end < start or (end - start).days >= 7 or end > datetime.now(MOSCOW).date():
        raise ValueError(_('Выберите период не более семи дней, без будущих дат.'))
    begin = datetime.combine(start, time.min, MOSCOW).astimezone(timezone.utc)
    until = datetime.combine(end + timedelta(days=1), time.min, MOSCOW).astimezone(timezone.utc)
    return start, end, begin, until

def group_from_form(data, current_id=None):
    name = data.get('new_group', '').strip()
    if name:
        if len(name) > 40:
            raise ValueError(_('Название группы: не более 40 символов.'))
        group = db.session.scalar(select(Group).where(Group.name == name))
        if not group:
            group = Group(name=name, active=True)
            group.admission_year = integer(data.get('admission_year') or group.entry_year, _('Год поступления'), 2000, 2100)
            db.session.add(group)
            db.session.flush()
            audit('catalog_updated', f'group:{group.id}', name)
    else:
        group = db.session.get(Group, integer(data.get('group_id'), _('Группа')))
    if not group or (not group.active and group.id != current_id):
        raise ValueError(_('Выберите действующую группу.'))
    return group

def department_from_form(data, current_id=None):
    name = data.get('new_department', '').strip()
    if name:
        if len(name) > 160:
            raise ValueError(_('Название кафедры: не более 160 символов.'))
        department = db.session.scalar(select(Department).where(Department.name == name))
        if not department:
            department = Department(name=name, active=True)
            db.session.add(department)
            db.session.flush()
            audit('catalog_updated', f'department:{department.id}', name)
    else:
        department = db.session.get(Department, integer(data.get('department_id'), _('Кафедра')))
    if not department or (not department.active and department.id != current_id):
        raise ValueError(_('Выберите действующую кафедру.'))
    return department

def form_catalogs():
    return dict(groups=db.session.scalars(select(Group).order_by(Group.name)).all(),
        departments=db.session.scalars(select(Department).order_by(Department.name)).all())


def validate_role_login(login, role, user_id=None, current_login=None):
    canonical = normalize_login(login)
    if role not in ('student', 'teacher', 'admin') or not re.fullmatch(r'[a-z0-9][a-z0-9._-]{1,119}', canonical):
        raise ValueError(_('Укажите корректные роль и логин.'))
    if role == 'student' and not canonical.isdigit():
        raise ValueError(_('Для студента логином должен быть номер студенческого из цифр.'))
    if role != 'student' and canonical.isdigit():
        raise ValueError(_('Для преподавателя или администратора нужен логин корпоративной почты.'))
    duplicate = db.session.scalar(select(User.id).where(
        User.login == canonical, User.role == role, User.id != user_id).limit(1))
    if duplicate:
        raise ValueError(_('Учётная запись с таким логином и ролью уже существует.'))
    if user_id is not None and canonical != current_login and db.session.scalar(select(User.id).where(
            User.login == canonical, User.id != user_id).limit(1)):
        raise ValueError(_('При изменении логина выберите значение, которое не используется другой учётной записью.'))
    return canonical


def submitted_admin_permissions(data, actor, target=None):
    requested = set(data.getlist('admin_permissions')) if hasattr(data, 'getlist') else set(data.get('admin_permissions', ()))
    if requested - ADMIN_PERMISSION_CODE_SET or requested - actor.admin_permission_codes:
        abort(403)
    preserved = target.admin_permission_codes - actor.admin_permission_codes if target and target.role == 'admin' else set()
    return requested | preserved


def set_admin_permissions(user, permissions):
    permissions = set(permissions)
    current = {grant.permission: grant for grant in user.permission_grants}
    for code in current.keys() - permissions:
        db.session.delete(current[code])
    for code in permissions - current.keys():
        user.permission_grants.append(AdminPermission(permission=code))


def can_manage_admin(actor, target):
    return bool(actor and target and actor.id != target.id and
        has_admin_permission(actor, 'admins.manage') and
        target.admin_permission_codes <= actor.admin_permission_codes)


def require_manage_user(target):
    if target.role == 'admin':
        if not can_manage_admin(g.user, target):
            abort(403)
    elif not has_admin_permission(g.user, 'users.manage'):
        abort(403)


def available_account_roles(actor):
    result = []
    if has_admin_permission(actor, 'users.manage'):
        result.extend(('student', 'teacher'))
    if has_admin_permission(actor, 'admins.manage'):
        result.append('admin')
    return tuple(result)


def ensure_admin_manager_remains(user, role=None, active=None, permissions=None):
    role = user.role if role is None else role
    active = user.active if active is None else active
    permissions = user.admin_permission_codes if permissions is None else set(permissions)
    removes_manager = (user.role == 'admin' and user.active and not user.deleted_at and
        'admins.manage' in user.admin_permission_codes and
        (role != 'admin' or not active or 'admins.manage' not in permissions))
    if not removes_manager:
        return
    another = db.session.scalar(select(AdminPermission.user_id).join(User).where(
        AdminPermission.permission == 'admins.manage', User.id != user.id,
        User.role == 'admin', User.active, User.deleted_at.is_(None)).limit(1))
    if not another:
        raise ValueError(_('Нельзя лишить полномочий последнего действующего администратора, который управляет правами.'))


def excel_text(value):
    value = str(value or '')
    return "'" + value if value.startswith(('=', '+', '-', '@')) else value


def user_status(user, gate):
    if user.deleted_at:
        return _('В корзине')
    if user.dismissed_on:
        return _('Уволен')
    if not user.active:
        return _('Отключён')
    if gate and gate.permanent:
        return _('Заблокирован')
    if gate and gate.locked_until and gate.locked_until > utcnow():
        return _('Временная блокировка')
    return _('Активен')


def local_stamp(value, date_only=False):
    if not value:
        return ''
    if isinstance(value, date) and not isinstance(value, datetime):
        return value.strftime('%d.%m.%Y')
    fmt = '%d.%m.%Y' if date_only else '%d.%m.%Y %H:%M'
    return value.astimezone(MOSCOW).strftime(fmt)


def add_user_sheet(book, role, users, gates):
    titles = {'student': _('Студенты'), 'teacher': _('Преподаватели'), 'admin': _('Администраторы')}
    sheet = book.create_sheet(str(titles[role])[:31])
    if role == 'student':
        headers = [_('ID'), _('Номер студенческого'), _('Фамилия'), _('Имя'), _('Отчество'), _('Группа'),
            _('Курс'), _('Форма обучения'), _('Профиль заполнен'), _('Статус'), _('Создан'), _('Удалён')]
        rows = [[u.id, excel_text(u.login), excel_text(u.last_name), excel_text(u.first_name), excel_text(u.middle_name),
            excel_text(u.group.name if u.group else ''), u.current_course if u.group else '',
            _('Заочная') if u.study_mode == 'part_time' else _('Очная') if u.study_mode == 'full_time' else '',
            _('Да') if u.name_locked and u.full_name and u.group else _('Нет'), user_status(u, gates.get(u.login)),
            local_stamp(u.created_at), local_stamp(u.deleted_at)] for u in users]
    elif role == 'teacher':
        headers = [_('ID'), _('Логин'), _('Фамилия'), _('Имя'), _('Отчество'), _('Кафедра'), _('Дисциплины'),
            _('Дата увольнения'), _('Статус'), _('Создан'), _('Удалён')]
        rows = [[u.id, excel_text(u.login), excel_text(u.last_name), excel_text(u.first_name), excel_text(u.middle_name),
            excel_text(u.department.name if u.department else ''), excel_text(', '.join(sorted(s.name for s in u.subjects))),
            local_stamp(u.dismissed_on, True), user_status(u, gates.get(u.login)), local_stamp(u.created_at),
            local_stamp(u.deleted_at)] for u in users]
    else:
        headers = [_('ID'), _('Логин'), _('Фамилия'), _('Имя'), _('Отчество'), _('Права администратора'), _('Статус'), _('Создан'), _('Удалён')]
        rows = [[u.id, excel_text(u.login), excel_text(u.last_name), excel_text(u.first_name), excel_text(u.middle_name),
            excel_text(', '.join(sorted(u.admin_permission_codes))), user_status(u, gates.get(u.login)),
            local_stamp(u.created_at), local_stamp(u.deleted_at)] for u in users]
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    fill = PatternFill('solid', fgColor='0646E8')
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(vertical='center')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 24
    for index, column in enumerate(sheet.columns, 1):
        width = max(len(str(cell.value or '')) for cell in column)
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 11), 42)
    return len(rows)


def student_logins_from_excel(upload):
    filename = (upload.filename or '').lower()
    if not filename.endswith('.xlsx'):
        raise StudentImportError(_('Выберите файл Excel в формате .xlsx.'))
    data = upload.read()
    if not data:
        raise StudentImportError(_('Файл Excel пуст.'))
    book = None
    try:
        with ZipFile(BytesIO(data)) as archive:
            if len(archive.infolist()) > 1000 or sum(item.file_size for item in archive.infolist()) > MAX_XLSX_UNCOMPRESSED:
                raise StudentImportError(_('Файл Excel слишком сложный или повреждён.'))
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        if len(book.worksheets) != 1:
            raise StudentImportError(_('В файле должен быть ровно один лист.'))
        sheet = book.active
        if sheet.max_row > MAX_STUDENT_IMPORT:
            raise StudentImportError(_('За один импорт допускается не более 1500 строк.'))
        if sheet.max_column > 1:
            raise StudentImportError(_('Номера студенческих должны находиться только в столбце A.'))
        logins = []
        for row_number, (value,) in enumerate(sheet.iter_rows(min_col=1, max_col=1, values_only=True), 1):
            if value is None or isinstance(value, str) and not value.strip():
                continue
            if isinstance(value, bool):
                login = ''
            elif isinstance(value, int):
                login = str(value)
            elif isinstance(value, float) and value.is_integer() and abs(value) < 10**15:
                login = str(int(value))
            else:
                login = str(value).strip()
            if not re.fullmatch(r'\d{2,120}', login):
                raise StudentImportError(_('Строка %(value0)s: номер студенческого должен содержать не менее двух цифр.', value0=row_number))
            logins.append(login)
        if not logins:
            raise StudentImportError(_('В столбце A нет номеров студенческих.'))
        return logins
    except StudentImportError:
        raise
    except (BadZipFile, InvalidFileException, KeyError, EOFError, OSError, ParseError,
            SyntaxError, UnicodeError, ValueError, TypeError, IndexError, OverflowError):
        raise StudentImportError(_('Не удалось прочитать файл Excel. Проверьте формат .xlsx.')) from None
    finally:
        if book is not None:
            book.close()

def live_user(user_id):
    user = db.get_or_404(User, user_id)
    if user.deleted_at:
        abort(404)
    return user

def transliterate_surname(surname):
    alphabet = dict(zip('абвгдеёжзийклмнопрстуфхцчшщъыьэюя', ['a','b','v','g','d','e','yo','zh','z','i','j','k','l','m','n','o','p','r','s','t','u','f','kh','ts','ch','sh','shch','','y','','e','yu','ya']))
    return ''.join(alphabet.get(c,c) for c in surname.lower()).capitalize()

def create_user(data, admin_permissions=None):
    role = data.get('role')
    login = validate_role_login(data.get('login', ''), role)
    gate = lock_gate(login)
    user = User(login=login, role=role, active=True, name_locked=role != 'student')
    if role == 'student':
        group = group_from_form(data)
        user.group = group
        user.course = integer(data.get('course'), _('Курс'), 1, 5)
        user.study_mode = data.get('study_mode','full_time')
        if user.study_mode not in ('full_time','part_time') or (user.course == 5 and user.study_mode != 'part_time'):
            raise ValueError(_('Пятый курс доступен только при заочной форме обучения.'))
        user.course = user.current_course
        password = 'Student'
    else:
        user.set_name_parts(*name_parts(data))
        password = (data.get('initial_password') or '').strip()
        if role == 'teacher' and not password:
            password = transliterate_surname(user.last_name)
        if role == 'teacher':
            user.department = department_from_form(data)
        if not password or len(password) > 128 or (role == 'admin' and len(password) < 8):
            raise ValueError(_('Укажите начальный пароль. Для администратора — от 8 до 128 символов.'))
    with db.session.no_autoflush:
        check_password_distinct(user, password)
    user.password_hash = hash_password(password)
    user.initial_password_hash = user.password_hash
    user.must_change_password = current_app.config['FORCE_INITIAL_PASSWORD_CHANGE']
    db.session.add(user)
    if role == 'admin':
        permissions = ADMIN_PERMISSION_CODE_SET if admin_permissions is None else set(admin_permissions)
        if permissions - ADMIN_PERMISSION_CODE_SET:
            raise ValueError(_('Выбран неизвестный набор прав администратора.'))
        set_admin_permissions(user, permissions)
    db.session.flush()
    audit('user_created', user.id, f'{user.role}; {user.login}')
    return user


def update_user_from_form(user, data):
    actor = g.user
    old_role, old_login = user.role, user.login
    old_permissions = set(user.admin_permission_codes)
    target_role = data.get('role', old_role)
    if target_role not in ('student', 'teacher', 'admin'):
        raise ValueError(_('Выберите существующую роль.'))

    if user.id == actor.id:
        if target_role != 'admin' or normalize_login(data.get('login', old_login)) != old_login:
            raise ValueError(_('Нельзя изменить собственную роль или логин из административной панели.'))
        permissions = set(user.admin_permission_codes)
    elif old_role == 'admin':
        if not can_manage_admin(actor, user):
            abort(403)
        permissions = submitted_admin_permissions(data, actor, user) if target_role == 'admin' else set()
    elif target_role == 'admin':
        if not has_admin_permission(actor, 'admins.manage'):
            abort(403)
        permissions = submitted_admin_permissions(data, actor)
    else:
        if not has_admin_permission(actor, 'users.manage'):
            abort(403)
        permissions = set()

    new_login = validate_role_login(data.get('login', old_login), target_role,
        user_id=user.id, current_login=old_login)
    reason = data.get('reason', '').strip()
    if (target_role != old_role or new_login != old_login) and not reason:
        raise ValueError(_('Укажите основание изменения роли или логина пользователя.'))
    ensure_admin_manager_remains(user, role=target_role, permissions=permissions)

    supplied_name = any(data.get(key, '').strip() for key in ('last_name', 'first_name', 'middle_name', 'full_name'))
    parts = name_parts(data) if supplied_name else ['', '', '']
    name = ' '.join(filter(None, parts))
    if supplied_name and (name != (user.full_name or '') or parts != [user.last_name or '', user.first_name or '', user.middle_name or '']):
        if not reason:
            raise ValueError(_('Укажите основание изменения ФИО или роли пользователя.'))
        old_name = user.full_name or 'Не заполнено'
        user.set_name_parts(*parts)
        user.name_locked = True
        audit('name_corrected', user.id, f'{old_name} → {user.full_name}; {reason[:300]}')
    if target_role != 'student' and not user.full_name:
        raise ValueError(_('Для преподавателя или администратора укажите фамилию и имя.'))

    group = None
    course = None
    mode = None
    department = None
    dismissed_on = None
    if target_role == 'student':
        group = group_from_form(data, user.group_id if old_role == 'student' else None)
        course = integer(data.get('course'), _('Курс'), 1, 5)
        mode = data.get('study_mode')
        if mode not in ('full_time', 'part_time') or (course == 5 and mode != 'part_time'):
            raise ValueError(_('Пятый курс доступен только заочникам.'))
        if group.entry_year:
            today = utcnow().astimezone(ZoneInfo(current_app.config['APP_TIMEZONE'])).date()
            academic_year = today.year - (today.month < 9)
            duration = 5 if mode == 'part_time' else 4
            course = max(1, min(academic_year - group.entry_year + 1, duration))
        if old_role == 'student':
            future = db.session.scalars(select(Booking).join(Event, Booking.event_id == Event.id).where(
                Booking.student_id == user.id, Booking.status == 'active', Booking.ends_at > utcnow())).all()
            if any(booking.event.allowed_course != course or
                    (booking.event.group_id and booking.event.group_id != group.id) for booking in future):
                raise ValueError(_('Новая группа или курс не подходят действующим записям. Сначала отмените их в журнале.'))
    elif target_role == 'teacher':
        department = department_from_form(data, user.department_id if old_role == 'teacher' else None)
        dismissed = data.get('dismissed_on', '').strip()
        dismissed_on = date.fromisoformat(dismissed) if dismissed else None
        if dismissed_on and dismissed_on > utcnow().astimezone(ZoneInfo(current_app.config['APP_TIMEZONE'])).date():
            raise ValueError(_('Дата увольнения не может быть в будущем.'))

    if old_role == 'student' and target_role != 'student':
        for booking_id in db.session.scalars(select(Booking.id).where(
                Booking.student_id == user.id, Booking.status == 'active', Booking.ends_at > utcnow())).all():
            cancel_booking(booking_id, actor, 'Изменение роли пользователя')
    if old_role == 'teacher' and (target_role != 'teacher' or (dismissed_on and not user.dismissed_on)):
        for event_id in db.session.scalars(select(Event.id).where(
                Event.teacher_id == user.id, Event.status == 'active', Event.ends_at > utcnow())).all():
            cancel_event(event_id, actor, 'Изменение роли или увольнение преподавателя')

    user.login = new_login
    user.role = target_role
    user.group = group
    user.group_id = group.id if group else None
    user.course = course
    user.study_mode = mode
    user.department = department
    user.department_id = department.id if department else None
    user.dismissed_on = dismissed_on
    if target_role != 'teacher':
        user.subjects.clear()
    if target_role == 'student':
        user.name_locked = bool(user.full_name)
    if target_role == 'teacher' and dismissed_on:
        user.active = False
    set_admin_permissions(user, permissions if target_role == 'admin' else set())

    role_changed = old_role != target_role
    permissions_changed = old_permissions != (permissions if target_role == 'admin' else set())
    if role_changed and target_role == 'admin':
        password = data.get('initial_password', '').strip()
        if not 8 <= len(password) <= 128:
            raise ValueError(_('При назначении администратора укажите начальный пароль от 8 до 128 символов.'))
        check_password_distinct(user, password)
        user.password_hash = hash_password(password)
        user.initial_password_hash = user.password_hash
        user.must_change_password = current_app.config['FORCE_INITIAL_PASSWORD_CHANGE']
        clear_gate(lock_gate(new_login))
    if role_changed or old_login != new_login or permissions_changed:
        user.session_version += 1
    if role_changed:
        user.tutorial_seen_at = None
        audit('user_role_changed', user.id, f'{old_role} → {target_role}; {reason[:300]}')
    if old_login != new_login:
        audit('user_login_changed', user.id, f'{old_login} → {new_login}; {reason[:300]}')
    if target_role == 'admin' and permissions_changed:
        audit('admin_permissions_updated', user.id, ','.join(sorted(permissions)) or 'none')
    audit('user_updated', user.id)

@bp.get('')
@bp.get('/')
@roles_required('admin')
def dashboard():
    counts = {role:db.session.scalar(select(func.count(User.id)).where(User.role == role, User.active)) for role in ('student','teacher','admin')}
    counts['events'] = db.session.scalar(select(func.count(Event.id)).where(Event.status == 'active', Event.ends_at > utcnow()))
    counts['trash'] = db.session.scalar(select(func.count(User.id)).where(User.deleted_at.is_not(None)))
    counts['graduates'] = sum(u.graduation_due for u in db.session.scalars(select(User).where(User.role == 'student', User.deleted_at.is_(None)).options(selectinload(User.group))).all())
    locked = db.session.scalars(select(LoginGate).where(or_(LoginGate.permanent,
        LoginGate.locked_until > utcnow())).order_by(LoginGate.updated_at.desc()).limit(20)).all() \
        if has_admin_permission(g.user, 'users.manage') else []
    scenario_available = {role: False for role in ('student', 'teacher')}
    if has_admin_permission(g.user, 'impersonation.use'):
        for role in scenario_available:
            scenario_query = select(User.id).where(
                User.role == role, User.active, User.deleted_at.is_(None))
            if role == 'teacher':
                scenario_query = scenario_query.where(User.dismissed_on.is_(None))
            scenario_available[role] = db.session.scalar(scenario_query.limit(1)) is not None
    logs = db.session.scalars(select(AuditLog).order_by(AuditLog.id.desc()).limit(8)).all() \
        if has_admin_permission(g.user, 'audit.view') else []
    return render_template('admin_dashboard.html', counts=counts, locked=locked, logs=logs,
        scenario_available=scenario_available)


def begin_impersonation(target):
    if not target or target.role not in ('student', 'teacher') or target.deleted_at:
        abort(404)
    if not target.active or (target.role == 'teacher' and target.dismissed_on):
        flash(_('Нельзя открыть режим для удалённой или отключённой учётной записи.'), 'error')
        return None
    admin = g.user
    session.update(admin_uid=admin.id, admin_version=admin.session_version,
        uid=target.id, version=target.session_version)
    session.permanent = True
    audit('admin_impersonation_started', target.id,
        f'{admin.login} -> {target.role}:{target.login}', actor_id=admin.id)
    db.session.commit()
    flash(_('Открыт режим работы от лица %(value0)s.', value0=target.full_name or target.login), 'success')
    return redirect(url_for('main.home'))


@bp.get('/impersonation')
@admin_permissions_required('impersonation.use')
def impersonation():
    query = select(User).where(User.role.in_(('student', 'teacher')), User.deleted_at.is_(None)).options(
        selectinload(User.group), selectinload(User.department)).order_by(User.role, User.full_name, User.login)
    role = request.args.get('role')
    if role in ('student', 'teacher'):
        query = query.where(User.role == role)
    term = request.args.get('q', '').strip()[:120]
    if term:
        query = query.where(or_(User.login.ilike('%' + term + '%'), User.full_name.ilike('%' + term + '%')))
    status = request.args.get('status', 'active')
    if status == 'active':
        query = query.where(User.active, or_(User.role != 'teacher', User.dismissed_on.is_(None)))
    elif status == 'inactive':
        query = query.where(or_(User.active.is_(False),
            and_(User.role == 'teacher', User.dismissed_on.is_not(None))))
    elif status != 'all':
        abort(400)
    page = db.paginate(query, per_page=50, max_per_page=50, error_out=False)
    gates = {user.login: db.session.get(LoginGate, user.login) for user in page.items}
    return render_template('admin_impersonation.html', page=page, gates=gates)

@bp.post('/impersonate/<role>')
@admin_permissions_required('impersonation.use')
def impersonate(role):
    if role not in ('student', 'teacher'):
        abort(404)
    target_query = select(User).where(
        User.role == role, User.active, User.deleted_at.is_(None))
    if role == 'teacher':
        target_query = target_query.where(User.dismissed_on.is_(None))
    target = db.session.scalar(target_query.order_by(User.id).limit(1))
    if not target:
        flash(_('Для выбранного сценария нет активных учётных записей.'), 'error')
        return redirect(url_for('admin.dashboard'))
    return begin_impersonation(target)


@bp.post('/impersonate/user/<int:user_id>')
@admin_permissions_required('impersonation.use')
def impersonate_user(user_id):
    target = db.session.get(User, user_id)
    result = begin_impersonation(target)
    if result:
        return result
    from .i18n import safe_return
    return redirect(safe_return(request.form.get('next') or url_for('admin.impersonation')))


@bp.post('/maintenance')
@admin_permissions_required('maintenance.manage')
def maintenance():
    setting = db.session.scalar(select(SiteSettings).where(SiteSettings.id == 1).with_for_update())
    enabled = request.form.get('maintenance') == 'on'
    if setting.maintenance != enabled:
        setting.maintenance = enabled
        audit('maintenance_enabled' if enabled else 'maintenance_disabled', 'site')
    db.session.commit()
    flash(_('Режим обслуживания включён.') if enabled else _('Режим обслуживания выключен.'), 'success')
    return redirect(url_for('admin.dashboard'))

@bp.post('/stop-impersonation')
def stop_impersonation():
    if not getattr(g, 'admin_user', None):
        abort(403)
    admin = g.admin_user
    session.clear()
    session.update(uid=admin.id, version=admin.session_version)
    session.permanent = True
    audit('admin_impersonation_stopped', admin.id, actor_id=admin.id)
    db.session.commit()
    flash(_('Режим работы от лица пользователя завершён.'), 'success')
    return redirect(url_for('admin.dashboard'))

@bp.get('/users')
@admin_permissions_required('users.manage', 'admins.manage', any_of=True)
def users():
    query = select(User).where(User.deleted_at.is_(None)).options(selectinload(User.group), selectinload(User.department)).order_by(User.role, User.login)
    if request.args.get('role') in ('student','teacher','admin'):
        query = query.where(User.role == request.args['role'])
    term = request.args.get('q','').strip()[:120]
    if term:
        query = query.where(or_(User.login.ilike('%'+term+'%'), User.full_name.ilike('%'+term+'%')))
    for arg, column in [('group', User.group_id), ('department', User.department_id)]:
        if request.args.get(arg, type=int):
            query = query.where(column == request.args.get(arg, type=int))
    if request.args.get('status') == 'graduated':
        marked = [u.id for u in db.session.scalars(query).all() if u.graduation_due]
        query = query.where(User.id.in_(marked))
    elif request.args.get('status') == 'dismissed':
        query = query.where(User.dismissed_on.is_not(None))
    elif request.args.get('status') == 'inactive':
        query = query.where(User.active.is_(False))
    page = db.paginate(query, per_page=50, max_per_page=50, error_out=False)
    gates = {u.login:db.session.get(LoginGate,u.login) for u in page.items}
    manageable_ids = {user.id for user in page.items if
        (user.role == 'admin' and can_manage_admin(g.user, user)) or
        (user.role != 'admin' and has_admin_permission(g.user, 'users.manage'))}
    editable_ids = manageable_ids | {g.user.id}
    if has_admin_permission(g.user, 'admins.manage'):
        editable_ids.update(user.id for user in page.items if user.role != 'admin')
    return render_template('admin_users.html', page=page, gates=gates,
        manageable_ids=manageable_ids, editable_ids=editable_ids,
        available_roles=available_account_roles(g.user), **form_catalogs())


@bp.post('/users/export')
@admin_permissions_required('users.export')
def export_users():
    scope = request.form.get('scope')
    if scope not in ('student', 'teacher', 'all'):
        abort(400)
    include_deleted = request.form.get('include_deleted') == 'yes'
    query = select(User).options(selectinload(User.group), selectinload(User.department), selectinload(User.subjects))
    if scope != 'all':
        query = query.where(User.role == scope)
    if not include_deleted:
        query = query.where(User.deleted_at.is_(None))
    exported = db.session.scalars(query.order_by(User.role, User.login, User.id)).all()
    gates = {gate.login: gate for gate in db.session.scalars(select(LoginGate).where(
        LoginGate.login.in_({user.login for user in exported}))).all()} if exported else {}
    book = Workbook()
    book.remove(book.active)
    roles_to_export = ('student', 'teacher', 'admin') if scope == 'all' else (scope,)
    counts = {}
    for role in roles_to_export:
        counts[role] = add_user_sheet(book, role, [user for user in exported if user.role == role], gates)
    book.properties.title = str(_('Выгрузка пользователей'))
    book.properties.creator = 'НФ НИТУ МИСИС'
    output = BytesIO()
    book.save(output)
    output.seek(0)
    audit('users_exported', scope, f'include_deleted={include_deleted}; counts={counts}')
    db.session.commit()
    exported_at = datetime.now(MOSCOW)
    return send_file(output, as_attachment=True,
        download_name=f'users-{scope}-{exported_at:%d-%m-%Y-%H-%M-%S}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', max_age=0)


@bp.route('/tutorials', methods=['GET', 'POST'])
@admin_permissions_required('tutorials.manage')
def tutorials():
    if request.method == 'POST':
        role = request.form.get('role')
        if role not in ('student', 'teacher', 'admin'):
            abort(400)
        action = request.form.get('action')
        try:
            video = db.session.get(TutorialVideo, role)
            if action == 'delete':
                if video:
                    db.session.delete(video)
                    clear_rutube_availability_cache()
                    audit('tutorial_deleted', role)
                db.session.commit()
                flash(_('Обучающее видео удалено. Пользователям будет показана заглушка.'), 'success')
            elif action == 'save':
                source_url = request.form.get('source_url', '').strip()
                embed_url = normalize_rutube_url(source_url)
                if not video:
                    video = TutorialVideo(role=role)
                    db.session.add(video)
                video.source_url = source_url
                video.embed_url = embed_url
                video.updated_at = utcnow()
                clear_rutube_availability_cache()
                audit('tutorial_updated', role)
                db.session.commit()
                flash(_('Ссылка на обучающее видео сохранена.'), 'success')
            else:
                abort(400)
            return redirect(url_for('admin.tutorials'))
        except ValueError as exc:
            db.session.rollback()
            flash(str(exc), 'error')
            return render_template('admin_tutorials.html', videos={
                item.role: item for item in db.session.scalars(select(TutorialVideo)).all()
            }), 400
    videos = {item.role: item for item in db.session.scalars(select(TutorialVideo)).all()}
    return render_template('admin_tutorials.html', videos=videos)

@bp.route('/users/new', methods=['GET','POST'])
@admin_permissions_required('users.manage', 'admins.manage', any_of=True)
def user_new():
    if request.method == 'POST':
        try:
            role = request.form.get('role')
            if role not in available_account_roles(g.user):
                abort(403)
            permissions = submitted_admin_permissions(request.form, g.user) if role == 'admin' else None
            user = create_user(request.form, admin_permissions=permissions)
            db.session.commit()
            flash(_('Учётная запись создана.'), 'success')
            return redirect(url_for('admin.user_edit', user_id=user.id))
        except ValueError as exc:
            db.session.rollback()
            flash(str(exc), 'error')
    return render_template('admin_user_form.html', user=None,
        available_roles=available_account_roles(g.user), grantable_permissions=g.user.admin_permission_codes,
        **form_catalogs())

@bp.route('/users/<int:user_id>', methods=['GET','POST'])
@admin_permissions_required('users.manage', 'admins.manage', any_of=True)
def user_edit(user_id):
    user = live_user(user_id)
    if user.role == 'admin' and user.id != g.user.id and not can_manage_admin(g.user, user):
        abort(403)
    if user.role != 'admin' and not (has_admin_permission(g.user, 'users.manage') or has_admin_permission(g.user, 'admins.manage')):
        abort(403)
    if request.method == 'POST':
        try:
            user = db.session.execute(select(User).where(User.id == user_id).with_for_update().execution_options(populate_existing=True)).scalar_one()
            if user.deleted_at:
                abort(404)
            update_user_from_form(user, request.form)
            db.session.commit()
            flash(_('Данные пользователя сохранены.'), 'success')
        except ValueError as exc:
            db.session.rollback()
            flash(str(exc), 'error')
        return redirect(url_for('admin.user_edit', user_id=user.id))
    return render_template('admin_user_form.html', user=user, gate=db.session.get(LoginGate,user.login),
        available_roles=available_account_roles(g.user), grantable_permissions=g.user.admin_permission_codes,
        target_manageable=(can_manage_admin(g.user, user) if user.role == 'admin'
            else has_admin_permission(g.user, 'users.manage')),
        **form_catalogs())

@bp.post('/users/<int:user_id>/reset')
@admin_permissions_required('users.manage', 'admins.manage', any_of=True)
def user_reset(user_id):
    user = live_user(user_id)
    require_manage_user(user)
    gate = lock_gate(user.login)
    user = db.session.execute(select(User).where(User.id == user_id).with_for_update().execution_options(populate_existing=True)).scalar_one()
    if user.deleted_at:
        abort(404)
    user.password_hash = user.initial_password_hash
    user.must_change_password = current_app.config['FORCE_INITIAL_PASSWORD_CHANGE']
    user.session_version += 1
    clear_gate(gate)
    # Logout all roles sharing the locked login; the other role's password is preserved.
    for other in db.session.scalars(select(User).where(User.login == user.login, User.id != user.id)).all():
        other.session_version += 1
    audit('password_reset',user.id,'Начальный пароль восстановлен; блокировка логина снята.')
    notify([user.id], _l('Пароль сброшен администратором'), _l('Восстановлен начальный пароль. При следующем входе смените его в настройках.'), link='/settings')
    db.session.commit()
    flash(_('Начальный пароль восстановлен. Блокировка общего логина снята.'), 'success')
    return redirect(url_for('admin.user_edit',user_id=user.id))

@bp.post('/users/<int:user_id>/toggle')
@admin_permissions_required('users.manage', 'admins.manage', any_of=True)
def user_toggle(user_id):
    user = live_user(user_id)
    if user.dismissed_on:
        flash(_('Для активации сначала снимите отметку об увольнении.'), 'error')
        return redirect(url_for('admin.user_edit', user_id=user.id))
    if user.id == g.user.id:
        flash(_('Нельзя отключить собственную учётную запись.'), 'error')
        return redirect(url_for('admin.user_edit',user_id=user.id))
    require_manage_user(user)
    ensure_admin_manager_remains(user, active=not user.active)
    lock_gate(user.login)
    user = db.session.execute(select(User).where(User.id == user_id).with_for_update().execution_options(populate_existing=True)).scalar_one()
    if user.deleted_at or user.dismissed_on:
        abort(409)
    if user.active:
        if user.role == 'teacher':
            for eid in db.session.scalars(select(Event.id).where(Event.teacher_id == user.id,Event.status == 'active',Event.ends_at > utcnow())).all():
                cancel_event(eid,g.user,'Отключение преподавателя администратором')
        if user.role == 'student':
            for bid in db.session.scalars(select(Booking.id).where(Booking.student_id == user.id,Booking.status == 'active',Booking.ends_at > utcnow())).all():
                cancel_booking(bid,g.user,'Отключение студента администратором')
    user.active = not user.active
    user.session_version += 1
    audit('user_activated' if user.active else 'user_deactivated',user.id)
    db.session.commit()
    flash(_('Статус учётной записи изменён.'), 'success')
    return redirect(url_for('admin.user_edit',user_id=user.id))

@bp.post('/users/<int:user_id>/delete')
@admin_permissions_required('users.manage', 'admins.manage', any_of=True)
def user_delete(user_id):
    try:
        user = live_user(user_id)
        if user.id == g.user.id:
            raise ValueError(_('Нельзя удалить собственную учётную запись.'))
        require_manage_user(user)
        ensure_admin_manager_remains(user, active=False)
        trash_users([user_id], g.user)
        db.session.commit()
        flash(_('Пользователь перенесён в корзину на 10 дней.'), 'success')
        return redirect(url_for('admin.users'))
    except ValueError as exc:
        db.session.rollback()
        flash(str(exc), 'error')
        return redirect(url_for('admin.users'))

@bp.post('/users/bulk')
@admin_permissions_required('users.manage', 'admins.manage', any_of=True)
def users_bulk():
    restoring = request.form.get('action') == 'restore'
    try:
        ids = {integer(value, _('Пользователь')) for value in request.form.getlist('user_ids')}
        if len(ids) > 500:
            raise ValueError(_('За один раз можно выбрать до 500 пользователей.'))
        selected = db.session.scalars(select(User).where(User.id.in_(ids))).all()
        for user in selected:
            if user.id != g.user.id:
                require_manage_user(user)
                if not restoring:
                    ensure_admin_manager_remains(user, active=False)
        count = restore_users(ids) if restoring else trash_users(ids, g.user)
        db.session.commit()
        flash(_('Восстановлено: %(value0)s.', value0=count) if restoring else _('Перенесено в корзину на 10 дней: %(value0)s.', value0=count), 'success')
    except ValueError as exc:
        db.session.rollback()
        flash(str(exc), 'error')
    return redirect(url_for('admin.trash' if restoring else 'admin.users'))

@bp.get('/trash')
@admin_permissions_required('users.manage', 'admins.manage', any_of=True)
def trash():
    query = select(User).where(User.deleted_at.is_not(None)).order_by(User.deleted_at.desc())
    page = db.paginate(query, per_page=50, error_out=False)
    manageable_ids = {user.id for user in page.items if
        (user.role == 'admin' and can_manage_admin(g.user, user)) or
        (user.role != 'admin' and has_admin_permission(g.user, 'users.manage'))}
    return render_template('admin_trash.html', page=page, manageable_ids=manageable_ids)

@bp.route('/catalogs', methods=['GET','POST'])
@admin_permissions_required('catalogs.manage')
def catalogs():
    if request.method == 'POST':
        try:
            kind = request.form.get('kind')
            if kind in ('group','subject','department'):
                model = {'group': Group, 'subject': Subject, 'department': Department}[kind]
                entity = db.get_or_404(model,integer(request.form.get('id'),'ID')) if request.form.get('id') else model()
                if kind == 'department' and request.form.get('delete') == 'yes':
                    if db.session.scalar(select(User.id).where(User.department_id == entity.id).limit(1)):
                        raise ValueError(_('Кафедра связана с пользователями. Перенесите их в другую кафедру или отправьте кафедру в архив.'))
                    audit('catalog_updated', f'department:{entity.id}', 'Кафедра удалена')
                    db.session.delete(entity)
                    db.session.commit()
                    flash(_('Кафедра удалена.'), 'success')
                    return redirect(url_for('admin.catalogs'))
                name = request.form.get('name','').strip()
                if not name or len(name) > (40 if kind == 'group' else 160 if kind == 'department' else 120):
                    raise ValueError(_('Название пустое или слишком длинное.'))
                entity.name = name
                if kind == 'group':
                    entity.admission_year = integer(request.form.get('admission_year') or entity.entry_year, _('Год поступления'), 2000, 2100)
                active = request.form.get('active') == 'yes'
                if entity.id and not active:
                    if kind == 'group' and db.session.scalar(select(User.id).where(User.group_id == entity.id, User.active).limit(1)):
                        raise ValueError(_('В группе есть действующие студенты. Сначала переведите или отключите их.'))
                    if kind == 'subject' and db.session.scalar(select(Event.id).where(Event.subject_id == entity.id,Event.status == 'active',Event.ends_at > utcnow()).limit(1)):
                        raise ValueError(_('По дисциплине есть будущие занятия. Сначала отмените их.'))
                entity.active = active
                db.session.add(entity)
                db.session.flush()
                audit('catalog_updated',f'{kind}:{entity.id}',name)
            elif kind == 'assignment':
                teacher = db.session.execute(select(User).where(User.id == integer(request.form.get('teacher_id'),_('Преподаватель'))).with_for_update()).scalar_one_or_none()
                subject = db.session.get(Subject,integer(request.form.get('subject_id'),_('Дисциплина')))
                if not teacher or teacher.role != 'teacher' or teacher.deleted_at or not subject:
                    raise ValueError(_('Выберите преподавателя и дисциплину.'))
                remove = request.form.get('remove') == 'yes'
                if remove:
                    if db.session.scalar(select(Event.id).where(Event.teacher_id == teacher.id,Event.subject_id == subject.id,Event.status == 'active',Event.ends_at > utcnow()).limit(1)):
                        raise ValueError(_('Сначала отмените будущие занятия по этой дисциплине.'))
                    if subject in teacher.subjects:
                        teacher.subjects.remove(subject)
                elif subject not in teacher.subjects:
                    teacher.subjects.append(subject)
                audit('assignment_removed' if remove else 'assignment_added',teacher.id,subject.name)
            else:
                abort(400)
            db.session.commit()
            flash(_('Справочник обновлён.'), 'success')
        except ValueError as exc:
            db.session.rollback()
            flash(str(exc), 'error')
        return redirect(url_for('admin.catalogs'))
    return render_template('admin_catalogs.html', **form_catalogs(), subjects=db.session.scalars(select(Subject).order_by(Subject.name)).all(), teachers=db.session.scalars(select(User).where(User.role == 'teacher', User.deleted_at.is_(None)).order_by(User.full_name)).all())

@bp.route('/import',methods=['GET','POST'])
@admin_permissions_required('users.import')
def import_students():
    if request.method == 'POST':
        try:
            file = request.files.get('file')
            if not file:
                raise ValueError(_('Выберите файл Excel.'))
            logins = student_logins_from_excel(file)
            unique_logins = list(dict.fromkeys(logins))
            existing = set(db.session.scalars(select(User.login).where(User.login.in_(unique_logins))).all())
            new_logins = [login for login in unique_logins if login not in existing]
            password_hashes = {prefix:hash_password('Student' + prefix) for prefix in {login[:2] for login in new_logins}}
            users = []
            for login in new_logins:
                password_hash = password_hashes[login[:2]]
                users.append(User(login=login, role='student', active=True, name_locked=False,
                    password_hash=password_hash, initial_password_hash=password_hash,
                    must_change_password=current_app.config['FORCE_INITIAL_PASSWORD_CHANGE']))
            if users:
                db.session.add_all(users)
                db.session.flush()
                db.session.execute(insert(LoginGate).values([
                    dict(login=user.login, failures=0, permanent=False, updated_at=utcnow()) for user in users
                ]).on_conflict_do_nothing())
            skipped = len(logins) - len(users)
            audit('students_imported','xlsx',f'rows={len(logins)}; created={len(users)}; skipped={skipped}')
            db.session.commit()
            flash(_('Добавлено студентов: %(value0)s. Пропущено повторов: %(value1)s.', value0=len(users), value1=skipped), 'success')
            return redirect(url_for('admin.users'))
        except ValueError as exc:
            db.session.rollback()
            flash(str(exc), 'error')
    return render_template('admin_import.html')

@bp.get('/events')
@admin_permissions_required('events.manage')
def events():
    query = select(Event).order_by(Event.starts_at.desc())
    page = db.paginate(query,per_page=20,max_per_page=20,error_out=False)
    return render_template('admin_events.html',page=page)

@bp.get('/audit')
@admin_permissions_required('audit.view')
def logs():
    query = select(AuditLog).options(selectinload(AuditLog.actor)).order_by(AuditLog.id.desc())
    if request.args.get('action'):
        query = query.where(AuditLog.action == request.args['action'][:60])
    page = db.paginate(query,per_page=30,max_per_page=30,error_out=False)
    today = datetime.now(MOSCOW).date()
    return render_template('admin_audit.html', page=page,
        lines=[audit_line(entry) for entry in page.items], today=today,
        week_start=today - timedelta(days=6))


@bp.post('/audit/export')
@admin_permissions_required('audit.view')
def export_logs():
    from .database_admin import confirm_password

    try:
        start, end, begin, until = audit_dates(request.form)
        confirm_password()
        output = tempfile.SpooledTemporaryFile(max_size=1024 * 1024, mode='w+b')
        try:
            exported_at = datetime.now(MOSCOW)
            output.write(('\ufeffИстория действий · НФ НИТУ МИСИС\n'
                f'Период: {start:%d.%m.%Y} — {end:%d.%m.%Y} (МСК)\n'
                f'Выгружено: {exported_at:%d.%m.%Y %H:%M:%S} (МСК)\n\n').encode('utf-8'))
            query = select(AuditLog).options(selectinload(AuditLog.actor)).where(
                AuditLog.created_at >= begin, AuditLog.created_at < until).order_by(AuditLog.created_at, AuditLog.id)
            count = 0
            for entry in db.session.scalars(query).yield_per(500):
                output.write((audit_line(entry) + '\n').encode('utf-8'))
                count += 1
            if not count:
                output.write('Записей за выбранный период нет.\n'.encode('utf-8'))
            output.seek(0)
            audit('audit_log_exported', 'audit_log', f'{start.isoformat()}..{end.isoformat()}; rows={count}')
            db.session.commit()
            response = send_file(output, mimetype='text/plain; charset=utf-8', as_attachment=True,
                download_name=f'logs-{exported_at:%d-%m-%Y-%H-%M-%S}.txt', max_age=0)
            response.call_on_close(output.close)
            return response
        except Exception:
            output.close()
            raise
    except ValueError as exc:
        db.session.rollback()
        flash(str(exc), 'error')
        return logs(), 400
