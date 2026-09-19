from flask_babel import gettext as _, lazy_gettext as _l
from datetime import date, datetime, time, timezone
from io import BytesIO
from zoneinfo import ZoneInfo
from flask import Blueprint, abort, current_app, flash, g, redirect, render_template, request, send_file, url_for
from sqlalchemy import select, or_
from sqlalchemy.orm import selectinload
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .models import db, User, Group, Subject, Event, Booking, utcnow
from .auth import roles_required, audit
from .services import book_event, cancel_booking, cancel_event, save_event, booking_problem

bp = Blueprint('main', __name__)

@bp.get('/')
@roles_required()
def home():
    return redirect(url_for({'student':'main.events','teacher':'main.teacher','admin':'admin.dashboard'}[g.user.role]))

@bp.get('/health')
def health():
    db.session.execute(select(1))
    return {'status':'ok'}


@bp.get('/service-status')
@roles_required()
def service_status():
    return {'maintenance': g.maintenance}


@bp.post('/event-view')
@roles_required('student', 'teacher')
def event_view():
    view = request.form.get('view')
    if view not in ('grid', 'list'):
        abort(400)
    g.user.event_view = view
    db.session.commit()
    from .i18n import safe_return
    return redirect(safe_return(request.form.get('next', '/')))

@bp.get('/events')
@roles_required('student')
def events():
    query = select(Event).where(Event.status == 'active', Event.ends_at > utcnow()).options(selectinload(Event.bookings)).order_by(Event.starts_at)
    if request.args.get('subject', type=int):
        query = query.where(Event.subject_id == request.args.get('subject', type=int))
    if request.args.get('teacher', type=int):
        query = query.where(Event.teacher_id == request.args.get('teacher', type=int))
    if request.args.get('mine','1') == '1':
        query = query.where(Event.allowed_course == g.user.current_course, or_(Event.group_id == None, Event.group_id == g.user.group_id))
    if request.args.get('date'):
        try:
            day = date.fromisoformat(request.args['date'])
            begin = datetime.combine(day, time.min, tzinfo=ZoneInfo(current_app.config['APP_TIMEZONE'])).astimezone(timezone.utc)
            query = query.where(Event.starts_at >= begin)
        except ValueError:
            flash(_('Некорректная дата фильтра.'), 'error')
    page = db.paginate(query, per_page=30 if g.user.event_view == 'list' else 12, max_per_page=30, error_out=False)
    mine = db.session.scalars(select(Booking).where(Booking.student_id == g.user.id, Booking.status == 'active', Booking.ends_at > utcnow())).all()
    problems = {}
    for event in page.items:
        problem = booking_problem(event, g.user)
        if not problem and any(b.starts_at < event.ends_at and b.ends_at > event.starts_at for b in mine):
            problem = _('Пересекается с вашей записью.')
        problems[event.id] = problem
    return render_template('events.html', page=page, problems=problems, subjects=db.session.scalars(select(Subject).where(Subject.active).order_by(Subject.name)).all(), teachers=db.session.scalars(select(User).where(User.role == 'teacher', User.active).order_by(User.full_name)).all(), mine_count=len(mine))

@bp.post('/events/<int:event_id>/book')
@roles_required('student')
def book(event_id):
    try:
        book_event(g.user.id, event_id)
        db.session.commit()
        flash(_('Вы записаны на занятие. Оно появилось в разделе «Мои записи».'), 'success')
    except ValueError as exc:
        db.session.rollback()
        flash(str(exc), 'error')
    return redirect(url_for('main.events'))

@bp.get('/bookings')
@roles_required('student')
def bookings():
    query = select(Booking).where(Booking.student_id == g.user.id).order_by(Booking.starts_at.desc())
    page = db.paginate(query, per_page=20, max_per_page=20, error_out=False)
    return render_template('bookings.html', page=page)

@bp.post('/bookings/<int:booking_id>/cancel')
@roles_required('student','admin')
def cancel(booking_id):
    try:
        cancel_booking(booking_id, g.user, request.form.get('reason','Отмена студентом'))
        db.session.commit()
        flash(_('Запись отменена.'), 'success')
    except ValueError as exc:
        db.session.rollback()
        flash(str(exc), 'error')
    return redirect(url_for('main.bookings' if g.user.role == 'student' else 'main.journal'))

@bp.get('/teacher')
@roles_required('teacher')
def teacher():
    query = select(Event).where(Event.teacher_id == g.user.id).order_by(Event.starts_at.desc())
    if request.args.get('subject', type=int):
        query = query.where(Event.subject_id == request.args.get('subject', type=int))
    page = db.paginate(query, per_page=30 if g.user.event_view == 'list' else 12, max_per_page=30, error_out=False)
    return render_template('teacher.html', page=page)

@bp.route('/events/new', methods=['GET','POST'])
@bp.route('/events/<int:event_id>/edit', methods=['GET','POST'])
@roles_required('teacher','admin')
def event_form(event_id=None):
    event = db.get_or_404(Event, event_id) if event_id else None
    if event and g.user.role == 'teacher' and event.teacher_id != g.user.id:
        abort(403)
    if request.method == 'POST':
        try:
            result = save_event(request.form, g.user, event_id)
            db.session.commit()
            flash(_('Занятие сохранено.'), 'success')
            return redirect(url_for('main.event_detail', event_id=result.id))
        except ValueError as exc:
            db.session.rollback()
            flash(str(exc), 'error')
    return render_template('event_form.html', event=event, groups=db.session.scalars(select(Group).where(Group.active).order_by(Group.name)).all(),
        teachers=db.session.scalars(select(User).where(User.role == 'teacher', User.active).order_by(User.full_name)).all(),
        subjects=db.session.scalars(select(Subject).where(Subject.active).order_by(Subject.name)).all() if g.user.role == 'admin' else [s for s in g.user.subjects if s.active])

@bp.get('/events/<int:event_id>')
@roles_required('student','teacher','admin')
def event_detail(event_id):
    event = db.get_or_404(Event, event_id)
    if g.user.role == 'student':
        problem = booking_problem(event, g.user)
        if not problem and db.session.scalar(select(Booking.id).where(Booking.student_id == g.user.id, Booking.status == 'active', Booking.starts_at < event.ends_at, Booking.ends_at > event.starts_at).limit(1)):
            problem = _('Пересекается с вашей записью.')
        return render_template('student_event.html', event=event, problem=problem)
    if g.user.role == 'teacher' and event.teacher_id != g.user.id:
        abort(403)
    query = select(Booking).outerjoin(User, Booking.student_id == User.id).where(Booking.event_id == event.id)
    if request.args.get('group', type=int):
        query = query.where(User.group_id == request.args.get('group', type=int))
    records = db.session.scalars(query.order_by(User.group_id, User.full_name)).all()
    return render_template('event_detail.html', event=event, records=records, groups=db.session.scalars(select(Group).order_by(Group.name)).all())

@bp.post('/events/<int:event_id>/cancel')
@roles_required('teacher','admin')
def cancel_class(event_id):
    try:
        cancel_event(event_id, g.user, request.form.get('reason',''))
        db.session.commit()
        flash(_('Занятие и все записи на него отменены.'), 'success')
    except ValueError as exc:
        db.session.rollback()
        flash(str(exc), 'error')
    return redirect(url_for('main.event_detail', event_id=event_id))

@bp.post('/attendance/<int:booking_id>')
@roles_required('teacher','admin')
def attendance(booking_id):
    booking = db.get_or_404(Booking, booking_id)
    if g.user.role == 'teacher' and booking.event.teacher_id != g.user.id:
        abort(403)
    value = request.form.get('attendance')
    if value not in ('pending','present','absent') or booking.status != 'active' or booking.starts_at > utcnow():
        abort(400)
    booking.attendance = value
    audit('attendance_changed', booking.id, value)
    db.session.commit()
    return redirect(url_for('main.event_detail', event_id=booking.event_id))

def journal_query():
    query = select(Booking).join(Event, Booking.event_id == Event.id).outerjoin(User, Booking.student_id == User.id)
    if g.user.role == 'teacher':
        query = query.where(Event.teacher_id == g.user.id)
    for arg, column in [('group', User.group_id), ('subject', Event.subject_id), ('event', Event.id), ('teacher',Event.teacher_id)]:
        if request.args.get(arg, type=int):
            query = query.where(column == request.args.get(arg, type=int))
    if request.args.get('status') in ('active','cancelled'):
        query = query.where(Booking.status == request.args['status'])
    return query.order_by(Event.starts_at.desc(), User.group_id, User.full_name)

@bp.get('/journal')
@roles_required('teacher','admin')
def journal():
    page = db.paginate(journal_query(), per_page=25, max_per_page=25, error_out=False)
    return render_template('journal.html', page=page, groups=db.session.scalars(select(Group).order_by(Group.name)).all(), subjects=db.session.scalars(select(Subject).order_by(Subject.name)).all() if g.user.role == 'admin' else g.user.subjects)

@bp.get('/journal.xlsx')
@roles_required('teacher','admin')
def export_journal():
    records = db.session.scalars(journal_query().limit(10001)).all()
    if len(records) > 10000:
        flash(_('Для выгрузки свыше 10 000 строк уточните фильтры.'), 'error')
        return redirect(url_for('main.journal'))
    book = Workbook()
    sheet = book.active
    sheet.title = 'Журнал консультаций'
    sheet.append(['ID занятия','Дисциплина','Преподаватель','Начало','Окончание','Аудитория','ФИО студента','Группа','Курс','Номер студенческого','Дата записи','Статус','Посещаемость'])
    statuses = {'active':'Активна','cancelled':'Отменена','pending':'Не отмечено','present':'Присутствовал','absent':'Не явился'}
    def dt(value):
        return value.astimezone(ZoneInfo(current_app.config['APP_TIMEZONE'])).replace(tzinfo=None)
    for row in records:
        e, s = row.event, row.student
        values = [e.id,e.subject.name,e.teacher.full_name if e.teacher else 'Удалённый преподаватель',dt(e.starts_at),dt(e.ends_at),e.room,s.full_name if s else 'Удалённый студент',s.group.name if s and s.group else '',s.current_course if s else '',s.login if s else '',dt(row.created_at),statuses[row.status],statuses[row.attendance]]
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value,str):
                cell.data_type = 's'  # Prevent spreadsheet formula injection.
            if isinstance(cell.value,datetime):
                cell.number_format = 'DD.MM.YYYY HH:MM'
    for cell in sheet[1]:
        cell.font = Font(bold=True,color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='234C5A')
    for col, width in zip('ABCDEFGHIJKLM',[12,34,38,22,22,20,38,16,10,24,22,18,20]):
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    book.save(output)
    output.seek(0)
    audit('journal_exported', 'xlsx', f'rows={len(records)}; filters={dict(request.args)}')
    db.session.commit()
    return send_file(output, as_attachment=True, download_name='consultations_journal.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
