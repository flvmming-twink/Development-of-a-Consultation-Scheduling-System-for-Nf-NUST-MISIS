from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select

from portal.models import AuditLog, db, TutorialVideo, User
from portal.tutorials import normalize_rutube_url
from conftest import sign_in


VIDEO_ID = '7716bd3e665725c3c008ae7ab4ff02e2'
VIDEO_URL = f'https://rutube.ru/video/{VIDEO_ID}/'


def rows(sheet):
    return list(sheet.iter_rows(values_only=True))


def test_admin_exports_students_teachers_and_all_users(app, client):
    sign_in(client, 'admin')

    response = client.post('/admin/users/export', data={'scope': 'student'})
    assert response.status_code == 200
    assert 'users-student-' in response.headers['Content-Disposition']
    book = load_workbook(BytesIO(response.data), read_only=True)
    assert book.sheetnames == ['Студенты']
    student_rows = rows(book['Студенты'])
    assert student_rows[0][:6] == ('ID', 'Номер студенческого', 'Фамилия', 'Имя', 'Отчество', 'Группа')
    assert any(row[1] == '2300431' and row[5] == 'БПИ-23' for row in student_rows[1:])
    assert all('парол' not in str(value).lower() for value in student_rows[0])

    response = client.post('/admin/users/export', data={'scope': 'teacher'})
    book = load_workbook(BytesIO(response.data), read_only=True)
    teacher_rows = rows(book['Преподаватели'])
    assert any(row[1] == 'kurenkov.ee' and row[5] == 'Информатика' and 'Базы данных' in row[6]
               for row in teacher_rows[1:])

    response = client.post('/admin/users/export', data={'scope': 'all'})
    book = load_workbook(BytesIO(response.data), read_only=True)
    assert book.sheetnames == ['Студенты', 'Преподаватели', 'Администраторы']
    assert any(row[1] == 'lxrdx' for row in rows(book['Администраторы'])[1:])


def test_user_export_is_admin_only_and_validates_scope(app, client):
    sign_in(client, 'student')
    assert client.post('/admin/users/export', data={'scope': 'all'}).status_code == 403
    client.post('/logout')
    sign_in(client, 'admin')
    assert client.post('/admin/users/export', data={'scope': 'unknown'}).status_code == 400


def test_user_export_can_include_trash(app, client):
    sign_in(client, 'admin')
    client.post(f'/admin/users/{app.config["IDS"]["student"]}/delete')
    current = load_workbook(BytesIO(client.post('/admin/users/export', data={'scope': 'student'}).data), read_only=True)
    assert all(row[1] != '2300431' for row in rows(current['Студенты'])[1:])
    archived = load_workbook(BytesIO(client.post('/admin/users/export', data={
        'scope': 'student', 'include_deleted': 'yes'
    }).data), read_only=True)
    assert any(row[1] == '2300431' and row[9] == 'В корзине' for row in rows(archived['Студенты'])[1:])


def test_rutube_url_normalization():
    assert normalize_rutube_url(VIDEO_URL) == f'https://rutube.ru/play/embed/{VIDEO_ID}'
    assert normalize_rutube_url(
        f'https://rutube.ru/video/private/{VIDEO_ID}/?p=Private_key-42'
    ) == f'https://rutube.ru/play/embed/{VIDEO_ID}/?p=Private_key-42'
    assert normalize_rutube_url(f'https://rutube.ru/play/embed/{VIDEO_ID}') == f'https://rutube.ru/play/embed/{VIDEO_ID}'


def test_admin_manages_role_tutorials(app, client):
    sign_in(client, 'admin')
    response = client.post('/admin/tutorials', data={
        'role': 'student', 'action': 'save', 'source_url': VIDEO_URL
    }, follow_redirects=True)
    assert response.status_code == 200
    assert 'Ссылка на обучающее видео сохранена.' in response.text
    with app.app_context():
        video = db.session.get(TutorialVideo, 'student')
        assert video.source_url == VIDEO_URL
        assert video.embed_url == f'https://rutube.ru/play/embed/{VIDEO_ID}'
        entry = db.session.scalar(select(AuditLog).where(
            AuditLog.action == 'tutorial_updated'
        ).order_by(AuditLog.id.desc()))
        assert VIDEO_URL not in (entry.details or '')

    response = client.post('/admin/tutorials', data={
        'role': 'teacher', 'action': 'save', 'source_url': ''
    })
    assert response.status_code == 400
    assert 'Укажите ссылку на видео RUTUBE.' in response.text
    response = client.post('/admin/tutorials', data={
        'role': 'teacher', 'action': 'save', 'source_url': f'https://example.org/video/{VIDEO_ID}'
    })
    assert response.status_code == 400

    response = client.post('/admin/tutorials', data={'role': 'student', 'action': 'delete'}, follow_redirects=True)
    assert 'Обучающее видео удалено.' in response.text
    with app.app_context():
        assert db.session.get(TutorialVideo, 'student') is None


def test_tutorial_is_shown_once_and_remains_in_settings(app, client, monkeypatch):
    with app.app_context():
        db.session.add(TutorialVideo(role='student', source_url=VIDEO_URL,
            embed_url=f'https://rutube.ru/play/embed/{VIDEO_ID}'))
        db.session.commit()
    monkeypatch.setattr('portal.tutorials.rutube_video_available', lambda url, referer: True)

    sign_in(client, 'student')
    page = client.get('/events')
    assert 'data-auto-open="1"' in page.text
    assert 'data-tutorial-open' not in page.text
    status = client.get('/tutorial/status')
    assert status.json == {'available': True, 'embed_url': f'https://rutube.ru/play/embed/{VIDEO_ID}'}

    response = client.post('/tutorial/dismiss', data={'next': '/events'})
    assert response.status_code == 302 and response.location.endswith('/events')
    with app.app_context():
        student = db.session.scalar(select(User).where(User.login == '2300431', User.role == 'student'))
        assert student.tutorial_seen_at is not None

    page = client.get('/settings')
    assert 'data-auto-open="0"' in page.text
    assert 'data-tutorial-open' in page.text
    assert 'Открыть видеоинструкцию' in page.text


def test_missing_tutorial_uses_unavailable_placeholder(app, client):
    sign_in(client, 'teacher')
    page = client.get('/teacher')
    assert 'data-auto-open="1"' in page.text
    assert 'data-status-url=""' in page.text
    assert 'Видео временно недоступно, попробуйте просмотреть его через настройки' in page.text
    assert client.get('/tutorial/status').json == {'available': False}
