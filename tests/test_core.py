import io
import re
from datetime import timedelta
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
import os
import pytest
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from openpyxl import load_workbook
from portal.models import db,User,Group,Event,Booking,LoginGate,AuditLog,utcnow
from portal.admin import create_user,transliterate_surname
from portal.auth import hash_password
from portal.services import book_event,booking_problem
from conftest import sign_in

@pytest.mark.parametrize('role',['student','teacher','admin'])
def test_role_dispatch(app,client,role):
    assert sign_in(client,role).status_code==302
    with client.session_transaction() as session:
        assert session['uid']==app.config['IDS'][role]
    assert client.get('/',follow_redirects=True).status_code==200

def test_email_alias_and_distinct_password(app,client):
    with app.app_context():
        create_user(dict(role='admin',login='kurenkov.ee',full_name='Администратор Проверки',initial_password='Shared!123'))
        db.session.commit()
    assert client.post('/login',data={'login':'KURENKOV.EE@MISIS.RU','password':'Kurenkov'}).status_code==302
    response=client.post('/settings',data={'action':'password','current_password':'Kurenkov','new_password':'Shared!123','repeat_password':'Shared!123'},follow_redirects=True)
    assert 'занят другой ролью' in response.text
    with app.app_context():
        assert db.session.get(User,app.config['IDS']['teacher']).password_hash

def test_ten_then_five_and_admin_reset(app,client):
    for i in range(9):
        assert client.post('/login',data={'login':'2300431','password':'wrong'}).status_code==401
    assert client.post('/login',data={'login':'2300431','password':'wrong'}).status_code==429
    assert sign_in(client).status_code==429
    with app.app_context():
        gate=db.session.get(LoginGate,'2300431')
        assert gate.failures==10
        gate.locked_until=utcnow()-timedelta(seconds=1)
        db.session.commit()
    for i in range(4):
        assert client.post('/login',data={'login':'2300431','password':'wrong'}).status_code==401
    assert client.post('/login',data={'login':'2300431','password':'wrong'}).status_code==423
    assert sign_in(client).status_code==423
    assert sign_in(client,'admin').status_code==302
    client.post(f'/admin/users/{app.config["IDS"]["student"]}/reset')
    client.post('/logout')
    assert sign_in(client).status_code==302

def test_empty_fields_and_unknown_login(app,client):
    for _ in range(11):
        assert client.post('/login',data={'login':'2300431','password':''}).status_code==400
    with app.app_context():
        assert db.session.get(LoginGate,'2300431').failures==0
    for _ in range(10):
        client.post('/login',data={'login':'unknown','password':'wrong'})
    assert client.post('/login',data={'login':'unknown','password':'wrong'}).status_code==429
    assert sign_in(client).status_code==302

def test_success_clears_attempts(app,client):
    for _ in range(3):
        client.post('/login',data={'login':'2300431','password':'wrong'})
    assert sign_in(client).status_code==302
    with app.app_context():
        assert db.session.get(LoginGate,'2300431').failures==0

def test_shared_login_lock_and_console_recovery(app,client):
    for _ in range(10):
        client.post('/login',data={'login':'kurenkov.ee','password':'wrong'})
    assert sign_in(client,'teacher').status_code==429
    assert sign_in(client,'admin').status_code==302
    client.post('/logout')
    result=app.test_cli_runner().invoke(args=['recover-admin','--login','lxrdx'],input='Recovery!123\nRecovery!123\n')
    assert result.exit_code==0,result.output
    assert client.post('/login',data={'login':'lxrdx','password':'Recovery!123'}).status_code==302

def test_permissions(app,client):
    sign_in(client)
    for path in ['/admin/','/teacher','/journal','/journal.xlsx','/events/new']:
        assert client.get(path).status_code==403,path
    client.post('/logout')
    sign_in(client,'teacher')
    assert client.get('/admin/users').status_code==403
    assert client.get(f'/events/{app.config["IDS"]["overlap"]}').status_code==403
    assert client.post(f'/events/{app.config["IDS"]["overlap"]}/cancel',data={'reason':'x'}).status_code==302
    with app.app_context():
        assert db.session.get(Event,app.config['IDS']['overlap']).status=='active'

def test_one_time_name_and_admin_correction(app,client):
    with app.app_context():
        u=create_user(dict(role='student',login='2300440',group_id=app.config['IDS']['group'],course=4,study_mode='full_time'))
        db.session.commit()
        uid=u.id
    client.post('/login',data={'login':'2300440','password':'Student'})
    assert client.post('/settings',data={'action':'name','full_name':'Петров Пётр Петрович','confirm_name':'yes'}).status_code==302
    assert client.post('/settings',data={'action':'name','full_name':'Иванов Иван Иванович','confirm_name':'yes'}).status_code==403
    client.post('/logout');sign_in(client,'admin')
    client.post(f'/admin/users/{uid}',data={'full_name':'Петров Пётр Петрович','group_id':app.config['IDS']['group'],'course':4,'study_mode':'full_time'})
    client.post(f'/admin/users/{uid}',data={'full_name':'Иванов Иван Иванович','reason':'Личное обращение','group_id':app.config['IDS']['group'],'course':4,'study_mode':'full_time'})
    with app.app_context():
        assert db.session.get(User,uid).full_name=='Иванов Иван Иванович'
        assert db.session.scalar(select(AuditLog.id).where(AuditLog.action=='name_corrected'))

def test_booking_overlap_duplicate_and_adjacent(app,client):
    sign_in(client)
    ids=app.config['IDS']
    client.post(f'/events/{ids["first"]}/book')
    duplicate=client.post(f'/events/{ids["first"]}/book',follow_redirects=True)
    assert 'уже записаны' in duplicate.text
    response=client.post(f'/events/{ids["overlap"]}/book',follow_redirects=True)
    assert 'уже есть занятие' in response.text
    client.post(f'/events/{ids["adjacent"]}/book')
    with app.app_context():
        bookings=db.session.scalars(select(Booking).where(Booking.status=='active')).all()
        assert len(bookings)==2

def test_deadline_course_and_group(app,client):
    ids=app.config['IDS'];sign_in(client)
    with app.app_context():
        e=db.session.get(Event,ids['first']);u=db.session.get(User,ids['student'])
        # Exact boundary: 48 hours allowed, one microsecond later rejected.
        boundary=e.starts_at-timedelta(hours=48)
        assert booking_problem(e,u,boundary) is None
        assert booking_problem(e,u,boundary+timedelta(microseconds=1))
        e.allowed_course=3;db.session.commit()
    response=client.post(f'/events/{ids["first"]}/book',follow_redirects=True)
    assert 'другого курса или группы' in response.text
    with app.app_context():
        e=db.session.get(Event,ids['first']);e.allowed_course=4;e.group_id=ids['other_group'];db.session.commit()
    response=client.post(f'/events/{ids["first"]}/book',follow_redirects=True)
    assert 'другого курса или группы' in response.text

def test_cancel_frees_capacity_and_rebook(app,client):
    ids=app.config['IDS'];sign_in(client)
    with app.app_context():
        e=db.session.get(Event,ids['first']);e.capacity=1;db.session.commit()
    client.post(f'/events/{ids["first"]}/book')
    with app.app_context():
        bid=db.session.scalar(select(Booking.id))
    client.post(f'/bookings/{bid}/cancel')
    client.post(f'/events/{ids["first"]}/book')
    with app.app_context():
        assert db.session.get(Booking,bid).status=='active'
        assert db.session.get(Event,ids['first']).occupied==1

def test_last_seat_full_and_closed_deadline(app,client):
    ids=app.config['IDS']
    with app.app_context():
        event=db.session.get(Event,ids['first']);event.capacity=1
        other=create_user(dict(role='student',login='2300455',group_id=ids['group'],course=4,study_mode='full_time'))
        other.full_name='Петров Пётр Петрович';other.name_locked=True
        db.session.commit()
    sign_in(client)
    client.post(f'/events/{ids["first"]}/book');client.post('/logout')
    client.post('/login',data={'login':'2300455','password':'Student'})
    response=client.post(f'/events/{ids["first"]}/book',follow_redirects=True)
    assert 'Свободных мест нет' in response.text
    with app.app_context():
        event=db.session.get(Event,ids['overlap'])
        event.starts_at=utcnow()+timedelta(hours=20)
        event.ends_at=event.starts_at+timedelta(hours=1)
        db.session.commit()
    response=client.post(f'/events/{ids["overlap"]}/book',follow_redirects=True)
    assert 'Запись закрывается за 48 часов' in response.text

def test_admin_pages_and_catalog_workflow(app,client):
    ids=app.config['IDS'];sign_in(client,'admin')
    paths=['/admin/','/admin/users','/admin/users/new','/admin/catalogs','/admin/events','/admin/import','/admin/audit','/journal','/events/new',f'/events/{ids["first"]}',f'/admin/users/{ids["student"]}']
    for path in paths:
        assert client.get(path).status_code==200,path
    client.post('/admin/catalogs',data={'kind':'group','name':'БПИ-26','active':'yes'})
    client.post('/admin/catalogs',data={'kind':'subject','name':'Теория вероятностей','active':'yes'})
    from portal.models import Subject
    with app.app_context():
        subject=db.session.scalar(select(Subject).where(Subject.name=='Теория вероятностей'));sid=subject.id
        assert db.session.scalar(select(Group.id).where(Group.name=='БПИ-26'))
    client.post('/admin/catalogs',data={'kind':'assignment','teacher_id':ids['teacher'],'subject_id':sid})
    with app.app_context():
        assert sid in [s.id for s in db.session.get(User,ids['teacher']).subjects]

def test_admin_impersonation_and_safe_delete(app,client):
    ids=app.config['IDS'];sign_in(client,'admin')
    assert client.post('/admin/impersonate/student').status_code==302
    student_page=client.get('/events')
    assert student_page.status_code==200
    assert 'Сценарий: Студент' in student_page.text
    assert client.post('/admin/stop-impersonation').status_code==302
    assert client.get('/admin/').status_code==200
    with app.app_context():
        spare=create_user(dict(role='admin',login='flvmming',full_name='Администратор flvmming',initial_password='Admin!1234'))
        db.session.commit()
        spare_id=spare.id
    assert client.post(f'/admin/users/{spare_id}/delete').status_code==302
    with app.app_context():
        assert db.session.get(User,spare_id).deleted_at is not None
    response=client.post(f'/admin/users/{ids["student"]}/delete',follow_redirects=True)
    assert 'перенесён в корзину' in response.text

def test_teacher_overlap_and_ownership(app,client):
    ids=app.config['IDS'];sign_in(client,'teacher')
    from zoneinfo import ZoneInfo
    with app.app_context():
        e=db.session.get(Event,ids['first'])
        tz=ZoneInfo(app.config['APP_TIMEZONE'])
        data=dict(subject_id=ids['other_subject'],starts_at=e.starts_at.astimezone(tz).strftime('%Y-%m-%dT%H:%M'),ends_at=e.ends_at.astimezone(tz).strftime('%Y-%m-%dT%H:%M'),room='301',capacity=10,allowed_course=4)
    response=client.post('/events/new',data=data)
    assert 'пересекающееся' in response.text
    data['subject_id']='9999'
    assert 'не закреплена' in client.post('/events/new',data=data).text

def test_cancel_event_propagates(app,client):
    ids=app.config['IDS'];sign_in(client)
    client.post(f'/events/{ids["first"]}/book');client.post('/logout');sign_in(client,'teacher')
    client.post(f'/events/{ids["first"]}/cancel',data={'reason':'Перенос'})
    with app.app_context():
        assert db.session.get(Event,ids['first']).status=='cancelled'
        assert db.session.scalar(select(Booking)).status=='cancelled'

def test_postgres_exclusion_constraints(app):
    if os.getenv('SKIP_NATIVE_CONSTRAINTS')=='1':
        pytest.skip('Native driver rollback check; SQL constraints are checked separately on the WASM stand.')
    ids=app.config['IDS']
    with app.app_context():
        first=db.session.get(Event,ids['first'])
        clone=Event(teacher_id=first.teacher_id,subject_id=first.subject_id,starts_at=first.starts_at,ends_at=first.ends_at,room='another',capacity=2,allowed_course=4)
        db.session.add(clone)
        with pytest.raises(IntegrityError):
            db.session.commit()
        db.session.rollback()
        book_event(ids['student'],ids['first']);db.session.commit()
        overlap=db.session.get(Event,ids['overlap'])
        db.session.add(Booking(event_id=overlap.id,student_id=ids['student'],starts_at=overlap.starts_at,ends_at=overlap.ends_at))
        with pytest.raises(IntegrityError):
            db.session.commit()
        db.session.rollback()

def test_export_scopes_students_and_excel_strings(app,client):
    ids=app.config['IDS'];sign_in(client)
    client.post(f'/events/{ids["first"]}/book');client.post('/logout');sign_in(client,'teacher')
    response=client.get('/journal.xlsx')
    assert response.status_code==200
    sheet=load_workbook(io.BytesIO(response.data)).active
    assert sheet.max_row==2
    assert sheet.cell(2,10).value=='2300431'
    client.post('/logout');sign_in(client,'second')
    sheet=load_workbook(io.BytesIO(client.get('/journal.xlsx').data)).active
    assert sheet.max_row==1

def test_csrf_and_session_revocation(app,client):
    app.config['WTF_CSRF_ENABLED']=True
    assert sign_in(client).status_code==400
    html=client.get('/login').text
    token=re.search(r'name="csrf_token" value="([^"]+)"',html).group(1)
    assert client.post('/login',data={'csrf_token':token,'login':'2300431','password':'Student'}).status_code==302
    with app.app_context():
        u=db.session.get(User,app.config['IDS']['student']);u.session_version+=1;db.session.commit()
    assert client.get('/bookings').status_code==302

def test_admin_creation_no_public_signup_and_import_atomic(app,client):
    assert client.get('/register').status_code==404
    sign_in(client,'admin')
    response=client.post('/admin/users/new',data={'role':'student','login':'2300450','group_id':app.config['IDS']['group'],'course':5,'study_mode':'full_time'})
    assert 'только при заочной' in response.text
    payload='login;group;course;study_mode\n2300450;БПИ-23;4;full_time\n2300431;БПИ-23;4;full_time\n'
    response=client.post('/admin/import',data={'file':(io.BytesIO(payload.encode()),'students.csv')},content_type='multipart/form-data')
    assert 'уже существует' in response.text
    with app.app_context():
        assert db.session.scalar(select(User.id).where(User.login=='2300450')) is None

def test_theme_password_and_transliteration(app,client):
    assert transliterate_surname('Бой')=='Boj'
    assert transliterate_surname('Табельская')=='Tabelskaya'
    sign_in(client)
    html=client.post('/settings',data={'action':'theme','theme':'dark'},follow_redirects=True).text
    assert 'data-theme="dark"' in html
    client.post('/settings',data={'action':'password','current_password':'Student','new_password':'StudentNew!','repeat_password':'StudentNew!'})
    client.post('/logout')
    assert sign_in(client).status_code==401
    assert client.post('/login',data={'login':'2300431','password':'StudentNew!'}).status_code==302

@pytest.mark.concurrency
def test_concurrent_last_seat(app):
    if os.getenv('SKIP_CONCURRENCY')=='1':
        pytest.skip('Requires native PostgreSQL, not a single-connection WASM database.')
    ids=app.config['IDS']
    with app.app_context():
        event=db.session.get(Event,ids['first']);event.capacity=1
        other=create_user(dict(role='student',login='2300477',group_id=ids['group'],course=4,study_mode='full_time'))
        other.full_name='Петров Пётр Петрович'
        db.session.commit();other_id=other.id
    barrier=Barrier(2)
    def attempt(uid):
        with app.app_context():
            barrier.wait(timeout=5)
            try:
                book_event(uid,ids['first']);db.session.commit();return True
            except ValueError:
                db.session.rollback();return False
    with ThreadPoolExecutor(max_workers=2) as pool:
        results=list(pool.map(attempt,[ids['student'],other_id]))
    assert sorted(results)==[False,True]
